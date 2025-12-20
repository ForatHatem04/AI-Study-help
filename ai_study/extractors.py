from pathlib import Path
from typing import Iterable

import pdfplumber
from docx import Document
from pptx import Presentation


def extract_text_from_pdf(path: Path) -> str:
    text_chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def extract_text_from_docx(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    return "\n".join(paragraphs).strip()


def extract_text_from_pptx(path: Path) -> str:
    presentation = Presentation(path)
    slides_text = []
    for slide in presentation.slides:
        slide_lines = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_lines.append(shape.text)
        slides_text.append("\n".join(slide_lines))
    return "\n".join(slides_text).strip()


def extract_text_from_files(files: Iterable[str]) -> str:
    paths = [Path(file) for file in files]
    content_chunks = []
    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            content = extract_text_from_pdf(path)
        elif suffix in {".doc", ".docx"}:
            content = extract_text_from_docx(path)
        elif suffix in {".ppt", ".pptx"}:
            content = extract_text_from_pptx(path)
        else:
            raise ValueError(f"Unsupported file type for {path}")
        content_chunks.append(content)
    return "\n\n".join(chunk for chunk in content_chunks if chunk)
